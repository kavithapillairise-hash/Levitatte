"""
Levitatte -> Rela: five-slide presentation (16:9), ivory / antique-gold light theme.

Same palette, typefaces and refined owl as the proposal document, so the deck and the paper read
as one system. Slides are hand-composed from rectangles and text boxes rather than placeholder
layouts, because the commercial slide needs column control the autolayouts will not give.

Two rules make the layout deterministic rather than eyeballed:

  1. Every block height is DERIVED by measure.py from the real font files.
  2. Leading is quoted ABSOLUTELY (a multiple of the point size) and converted to the renderer's
     percentage-of-natural-line-height at write time, so a measured height and a rendered height
     are the same number by construction.

guard() then refuses to save a file containing a shape that is off-slide, non-positive, or
across the footer rule.

    python3 deck.py  Levitatte-Rela-5-Slide-Presentation.pptx
"""
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

import content as C
import measure as ms

HERE = os.path.dirname(os.path.abspath(__file__))
OWL = os.path.join(HERE, "owl_refined.png")

# ── palette, lifted verbatim from the proposal renderer ─────────────────────────
INK = RGBColor(0x20, 0x1C, 0x15)
BODY = RGBColor(0x55, 0x51, 0x49)
MUTE = RGBColor(0x8A, 0x86, 0x7C)
FAINT = RGBColor(0xB1, 0xAC, 0xA1)
GOLD = RGBColor(0x93, 0x74, 0x25)
GOLD_L = RGBColor(0xB0, 0x89, 0x30)
TAGCOL = RGBColor(0x4A, 0x42, 0x33)
PAPER = RGBColor(0xFA, 0xF8, 0xF3)
LINE = RGBColor(0xDD, 0xD7, 0xC8)
CARD = RGBColor(0xFC, 0xF8, 0xEF)
HEAD = RGBColor(0xF1, 0xED, 0xE2)
ZEBRA = RGBColor(0xFC, 0xFA, 0xF5)
RULE_DK = RGBColor(0x3A, 0x34, 0x28)
ON_DARK = RGBColor(0xEC, 0xE7, 0xDA)
ON_DARK_MUTE = RGBColor(0xA9, 0xA2, 0x92)

DISP = "Playfair Display"
SANS = "Inter"
BULLET = "–"

# absolute leading, as a multiple of the point size
LEAD_BODY = 1.36
LEAD_TIGHT = 1.30
LEAD_LABEL = 1.18
LEAD_ROW = 1.14
LEAD_DISP = 1.04

W, H = 13.333, 7.5
M = 0.62                       # side margin
CW = W - 2 * M                 # content width, 12.093
FOOT_Y = 7.03                  # footer rule
FLOOR = FOOT_Y - 0.16          # lowest a content block may reach

_SHAPES = []


# ── primitives ─────────────────────────────────────────────────────────────────
def _reg(sl, tag, x, y, w, h):
    _SHAPES.append((getattr(sl, "_tc_page", "?"), tag, x, y, w, h))


def rect(sl, x, y, w, h, fill=None, line=None, lw=0.5, tag="rect"):
    _reg(sl, tag, x, y, w, h)
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    # python-pptx attaches <p:style> referencing the theme's fill/line/effect/font. LibreOffice
    # honours its effectRef, which is why the first build had a drop shadow on every card and
    # zebra row. Drop the element and rely on the explicit formatting below.
    el = s._element
    st = el.find(qn("p:style"))
    if st is not None:
        el.remove(st)
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    return s


def hair(sl, x, y, w, color=LINE, h=0.0075):
    return rect(sl, x, y, w, h, fill=color, tag="hair")


def tbox(sl, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True, tag="text"):
    _reg(sl, tag, x, y, w, h)
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def para(tf, align=PP_ALIGN.LEFT, before=0, after=0, lead=None, f=SANS, b=False, i=False):
    """Reuse paragraph 0 while it is still empty, then append.

    Deliberately stateless: an earlier version cached "have I seen this frame" against id(tf),
    and CPython recycles ids once a TextFrame is collected, so unrelated frames collided and
    inherited an empty leading paragraph that pushed their text down by a line.
    """
    p0 = tf.paragraphs[0]
    p = p0 if (len(tf.paragraphs) == 1 and not p0.runs) else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    if lead:
        p.line_spacing = ms.spacing_pct(lead, f, b, i)
    return p


def run(p, text, f=SANS, sz=9, c=BODY, b=False, i=False, track=None):
    r = p.add_run()
    r.text = text
    r.font.name = f
    r.font.size = Pt(sz)
    r.font.bold = b
    r.font.italic = i
    r.font.color.rgb = c
    if track:
        r.font._rPr.set("spc", str(int(round(track * 100))))
    return r


def slide(prs, page, bg=PAPER):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    sl._tc_page = page
    rect(sl, 0, 0, W, H, fill=bg, tag="bg")
    return sl


def h_of(text, w, sz, *, f=SANS, b=False, i=False, track=0.0, lead=LEAD_BODY, pad=0.035):
    return ms.height_in(text, w, sz, f, b, i, track, lead, pad)


def distribute(top, bottom, heights, *, min_gap=0.05, tag=""):
    """Lay `heights` between top and bottom, spreading the slack as equal gaps.

    Raises rather than silently overlapping, which is how the first build ended up drawing a
    card on top of a bullet list.
    """
    need = sum(heights)
    n = len(heights)
    slack = (bottom - top) - need
    if slack < min_gap * (n - 1) - 1e-6:
        raise SystemExit(
            f"LAYOUT OVERFLOW in {tag or 'block'}: {n} items need {need:.3f}in plus "
            f"{min_gap * (n - 1):.3f}in of gaps, but only {bottom - top:.3f}in is available "
            f"(short by {min_gap * (n - 1) - slack:.3f}in)")
    gap = slack / (n - 1) if n > 1 else 0.0
    ys, y = [], top
    for h in heights:
        ys.append(y)
        y += h + gap
    return ys, gap


# ── composed elements ──────────────────────────────────────────────────────────
def mark(sl, x, y, scale=1.0, dark=False):
    ow = 0.46 * scale
    sl.shapes.add_picture(OWL, Inches(x), Inches(y), width=Inches(ow))
    tf = tbox(sl, x + ow + 0.10, y + 0.02 * scale, 1.6 * scale, 0.40 * scale)
    p = para(tf, lead=1.0, f=DISP, b=True)
    run(p, "LEVITATTE", DISP, 10.5 * scale, ON_DARK if dark else INK, b=True, track=1.5)
    p = para(tf, lead=LEAD_LABEL, b=True)
    run(p, "LEARNING & DEVELOPMENT", SANS, 4.6 * scale, GOLD_L if dark else GOLD, b=True,
        track=1.05)


def header(sl, num, eyebrow, title, dek=None):
    tf = tbox(sl, M, 0.44, CW - 2.5, 0.22)
    p = para(tf, lead=LEAD_LABEL)
    run(p, f"{num:02d}   ", DISP, 9, GOLD_L, b=True)
    run(p, eyebrow.upper(), SANS, 7.2, GOLD, b=True, track=2.3)

    tw = CW - 2.6
    tf = tbox(sl, M, 0.66, tw, 0.48)
    p = para(tf, lead=LEAD_DISP, f=DISP, b=True)
    run(p, title, DISP, 25, INK, b=True)

    y = 1.22
    if dek:
        dh = h_of(dek, tw, 10, f=DISP, i=True, lead=1.30, pad=0.04)
        tf = tbox(sl, M, 1.20, tw, dh)
        p = para(tf, lead=1.30, f=DISP, i=True)
        run(p, dek, DISP, 10, TAGCOL, i=True)
        y = 1.20 + dh + 0.10
    hair(sl, M, y, CW, color=GOLD, h=0.011)
    mark(sl, W - M - 2.06, 0.42)
    return y + 0.20


def footer(sl, page):
    hair(sl, M, FOOT_Y, CW)
    tf = tbox(sl, M, FOOT_Y + 0.11, CW - 0.6, 0.18)
    p = para(tf, lead=LEAD_LABEL)
    run(p, C.META["footer"], SANS, 6.2, MUTE, track=0.6)
    tf = tbox(sl, W - M - 0.6, FOOT_Y + 0.11, 0.6, 0.18)
    p = para(tf, align=PP_ALIGN.RIGHT, lead=LEAD_LABEL)
    run(p, str(page), SANS, 6.2, MUTE, b=True)


def block_label(sl, x, y, w, text, rule=True, sz=6.6):
    tf = tbox(sl, x, y, w, 0.17)
    p = para(tf, lead=LEAD_LABEL)
    run(p, text.upper(), SANS, sz, GOLD, b=True, track=2.1)
    if rule:
        hair(sl, x, y + 0.19, w)
    return y + 0.30


def body(sl, x, y, w, text, *, sz=7.8, lead=LEAD_BODY, color=BODY, italic=False,
         align=PP_ALIGN.JUSTIFY, f=SANS, bold=False, pad=0.035):
    h = h_of(text, w, sz, f=f, b=bold, i=italic, lead=lead, pad=pad)
    tf = tbox(sl, x, y, w, h)
    p = para(tf, align=align, lead=lead, f=f, b=bold, i=italic)
    run(p, text, f, sz, color, b=bold, i=italic)
    return y + h


def fit_blocks(top, bottom, height_fn, sizes, *, min_gap=0.05, tag=""):
    """Largest size in `sizes` whose measured blocks fit between top and bottom.

    Keeps type as large as the space honestly allows instead of hard-coding a size that
    happens to work today and silently overflows when a line of copy changes.
    """
    for sz in sizes:
        hs = height_fn(sz)
        if sum(hs) + min_gap * (len(hs) - 1) <= (bottom - top) + 1e-6:
            ys, _ = distribute(top, bottom, hs, min_gap=min_gap, tag=tag)
            return sz, ys, hs
    raise SystemExit(f"LAYOUT OVERFLOW in {tag or 'block'}: even {sizes[-1]}pt will not fit "
                     f"in {bottom - top:.3f}in")


def bullet_heights(items, w, sz, lead=LEAD_BODY, indent=0.17):
    return [h_of(it, w - indent, sz, lead=lead, pad=0.02) for it in items]


def bullets_at(sl, x, ys, w, items, *, sz=7.6, lead=LEAD_BODY, indent=0.17,
               color=BODY, lead_color=INK, lead_bold=True):
    for it, ry in zip(items, ys):
        head, sep, tail = it.partition(": ")
        tf = tbox(sl, x, ry, 0.16, 0.20)
        p = para(tf, lead=lead)
        run(p, BULLET, SANS, sz, GOLD_L, b=True)
        tf = tbox(sl, x + indent, ry, w - indent,
                  h_of(it, w - indent, sz, lead=lead, pad=0.02))
        p = para(tf, lead=lead)
        if sep and lead_bold and len(head) <= 62:
            run(p, head + ":", SANS, sz, lead_color, b=True)
            run(p, " " + tail, SANS, sz, color)
        else:
            run(p, it, SANS, sz, color)


def bullets(sl, x, y, w, items, *, sz=7.6, gap=0.11, lead=LEAD_BODY, indent=0.17,
            color=BODY, lead_color=INK, lead_bold=True):
    hs = bullet_heights(items, w, sz, lead, indent)
    ys, ry = [], y
    for h in hs:
        ys.append(ry)
        ry += h + gap
    bullets_at(sl, x, ys, w, items, sz=sz, lead=lead, indent=indent, color=color,
               lead_color=lead_color, lead_bold=lead_bold)
    return ry - gap


def table(sl, x, y, widths, header_row, rows, *, row_h, head_h=0.26,
          hsz=6.3, bsz=7.0, zebra=True, total=False, aligns=None,
          col_sizes=None, col_colors=None, col_bold=None):
    aligns = aligns or [PP_ALIGN.LEFT] * len(widths)
    col_sizes = col_sizes or [bsz] * len(widths)
    col_colors = col_colors or [INK if i == 0 else BODY for i in range(len(widths))]
    col_bold = col_bold or [i == 0 for i in range(len(widths))]
    tw = sum(widths)

    rect(sl, x, y, tw, head_h, fill=HEAD, tag="thead")
    cx = x
    for i, w in enumerate(widths):
        tf = tbox(sl, cx + 0.09, y, w - 0.18, head_h, anchor=MSO_ANCHOR.MIDDLE)
        p = para(tf, align=aligns[i], lead=LEAD_LABEL)
        run(p, header_row[i].upper(), SANS, hsz, GOLD, b=True, track=1.6)
        cx += w
    ry = y + head_h

    n = len(rows)
    for r_i, row in enumerate(rows):
        is_total = total and r_i == n - 1
        h = row_h + (0.02 if is_total else 0)
        if is_total:
            rect(sl, x, ry, tw, h, fill=HEAD, tag="trow")
        elif zebra and r_i % 2 == 1:
            rect(sl, x, ry, tw, h, fill=ZEBRA, tag="trow")
        hair(sl, x, ry, tw, h=0.006)
        cx = x
        for i, w in enumerate(widths):
            tf = tbox(sl, cx + 0.09, ry, w - 0.18, h, anchor=MSO_ANCHOR.MIDDLE)
            p = para(tf, align=aligns[i], lead=LEAD_ROW)
            run(p, str(row[i]), SANS, col_sizes[i],
                INK if is_total else col_colors[i], b=True if is_total else col_bold[i])
            cx += w
        ry += h
    hair(sl, x, ry, tw, color=GOLD, h=0.009)
    return ry + 0.02


def kv(sl, x, y, w, rows, *, label_w, lsz=5.9, vsz=7.0, pad=0.12, min_h=0.28):
    ry = y
    for i, (k, v) in enumerate(rows):
        lh = h_of(k, label_w - 0.12, lsz, b=True, track=1.4, lead=LEAD_LABEL, pad=0)
        vh = h_of(v, w - label_w, vsz, lead=LEAD_ROW, pad=0)
        h = max(lh, vh, min_h - pad) + pad
        if i:
            hair(sl, x, ry, w, h=0.006)
        tf = tbox(sl, x, ry, label_w - 0.12, h, anchor=MSO_ANCHOR.MIDDLE)
        p = para(tf, lead=LEAD_LABEL)
        run(p, k, SANS, lsz, GOLD, b=True, track=1.4)
        tf = tbox(sl, x + label_w, ry, w - label_w, h, anchor=MSO_ANCHOR.MIDDLE)
        p = para(tf, lead=LEAD_ROW)
        run(p, v, SANS, vsz, INK)
        ry += h
    return ry


def note_card(sl, x, y, w, text, label="NOTE", sz=6.8, lsz=6.2):
    inner = w - 0.32
    h = h_of(label + "     " + text, inner, sz, i=True, lead=LEAD_TIGHT) + 0.28
    rect(sl, x, y, w, h, fill=CARD, line=LINE, lw=0.5, tag="note")
    tf = tbox(sl, x + 0.16, y + 0.14, inner, h - 0.28)
    p = para(tf, lead=LEAD_TIGHT)
    run(p, label + "   ", SANS, lsz, GOLD, b=True, track=1.6)
    run(p, text, SANS, sz, BODY, i=True)
    return y + h


def accent_card(sl, x, y, w, text, *, sz=7.8, color=TAGCOL, italic=True, pad=0.32):
    inner = w - 0.40
    h = h_of(text, inner, sz, i=italic, lead=LEAD_TIGHT) + pad
    rect(sl, x, y, w, h, fill=CARD, line=LINE, lw=0.5, tag="accent")
    rect(sl, x, y, 0.028, h, fill=GOLD, tag="accentrule")
    tf = tbox(sl, x + 0.20, y, inner, h, anchor=MSO_ANCHOR.MIDDLE)
    p = para(tf, lead=LEAD_TIGHT, i=italic)
    run(p, text, SANS, sz, color, i=italic)
    return y + h


# ═══════════════════════════════════════════════════════════════════════════════
def slide1(prs):
    sl = slide(prs, 1)

    px, pw = 8.86, 3.85
    rect(sl, px, 0, pw + M, H, fill=INK, tag="panel")
    inner = pw - 0.84
    tf = tbox(sl, px + 0.42, 1.00, inner, 0.20)
    p = para(tf, lead=LEAD_LABEL)
    run(p, "THE ENGAGEMENT", SANS, 6.6, GOLD_L, b=True, track=2.2)
    hair(sl, px + 0.42, 1.26, inner, color=GOLD)

    stats = [("90", "DAYS"), ("39", "SESSIONS"), ("156", "FACILITATOR-LED HOURS"),
             ("12", "CUSTOMISED MODULES"), ("30", "PARTICIPANTS PER BATCH")]
    sy, step = 1.46, 0.70
    for i, (num, lab) in enumerate(stats):
        if i:
            hair(sl, px + 0.42, sy - 0.13, inner, color=RULE_DK, h=0.006)
        tf = tbox(sl, px + 0.42, sy, 1.05, 0.40)
        p = para(tf, lead=LEAD_DISP, f=DISP, b=True)
        run(p, num, DISP, 25, GOLD_L, b=True)
        tf = tbox(sl, px + 1.52, sy + 0.10, inner - 1.10, 0.32)
        p = para(tf, lead=LEAD_LABEL)
        run(p, lab, SANS, 6.3, ON_DARK_MUTE, b=True, track=1.5)
        sy += step

    cav = ("Sessions and hours are upper bounds: 36 to 39 sessions and 144 to 156 hours, "
           "subject to start date and hospital holidays.")
    ch = h_of(cav, inner, 6.2, i=True, lead=LEAD_TIGHT)
    tf = tbox(sl, px + 0.42, sy - 0.09, inner, ch)
    p = para(tf, lead=LEAD_TIGHT, i=True)
    run(p, cav, SANS, 6.2, ON_DARK_MUTE, i=True)

    fy = sy - 0.09 + ch + 0.34
    hair(sl, px + 0.42, fy, inner, color=GOLD)
    tf = tbox(sl, px + 0.42, fy + 0.22, inner, 0.44)
    p = para(tf, lead=LEAD_DISP, f=DISP, b=True)
    run(p, C.FEE_HEADLINE, DISP, 24, ON_DARK, b=True)
    sub_y = fy + 0.74
    hh = h_of(C.FEE_HEADLINE_SUB, inner, 6.9, lead=LEAD_TIGHT)
    tf = tbox(sl, px + 0.42, sub_y, inner, hh)
    p = para(tf, lead=LEAD_TIGHT)
    run(p, C.FEE_HEADLINE_SUB, SANS, 6.9, ON_DARK_MUTE)
    tf = tbox(sl, px + 0.42, sub_y + hh + 0.07, inner, 0.20)
    p = para(tf, lead=LEAD_TIGHT)
    run(p, C.FEE_TOTAL_LINE, SANS, 6.9, GOLD_L, b=True)

    # left composition
    sl.shapes.add_picture(OWL, Inches(M + 0.02), Inches(0.62), width=Inches(1.16))
    tf = tbox(sl, M, 1.48, 5.0, 0.34)
    p = para(tf, lead=LEAD_DISP, f=DISP, b=True)
    run(p, "LEVITATTE", DISP, 25, INK, b=True, track=4.2)
    tf = tbox(sl, M, 1.90, 5.0, 0.20)
    p = para(tf, lead=LEAD_LABEL)
    run(p, "LEARNING  &  DEVELOPMENT", SANS, 7.0, GOLD, b=True, track=3.0)

    hair(sl, M, 2.42, 2.0, color=GOLD, h=0.014)

    ty = 2.68
    for ln in C.META["deck_title_lines"]:
        tf = tbox(sl, M, ty, 7.8, 0.66)
        p = para(tf, lead=LEAD_DISP, f=DISP, b=True)
        run(p, ln, DISP, 43, INK, b=True)
        ty += 0.68

    tf = tbox(sl, M, ty + 0.14, 7.6, 0.30)
    p = para(tf, lead=1.22, f=DISP, i=True)
    run(p, C.META["subtitle"], DISP, 14, TAGCOL, i=True)
    tf = tbox(sl, M, ty + 0.54, 7.6, 0.20)
    p = para(tf, lead=LEAD_LABEL)
    run(p, C.META["descriptor"], SANS, 6.6, GOLD, b=True, track=2.0)

    sy2 = ty + 0.92
    hair(sl, M, sy2, 7.6)
    tf = tbox(sl, M, sy2 + 0.16, 6.0, 0.22)
    p = para(tf, lead=LEAD_ROW)
    run(p, C.META["prepared_for"], SANS, 8.2, INK)
    tf = tbox(sl, M + 6.1, sy2 + 0.17, 1.5, 0.20)
    p = para(tf, align=PP_ALIGN.RIGHT, lead=LEAD_LABEL)
    run(p, C.META["date"], SANS, 7.4, MUTE, b=True, track=1.6)
    hair(sl, M, sy2 + 0.50, 7.6)

    tf = tbox(sl, M, sy2 + 0.74, 7.7, 0.30)
    p = para(tf, lead=1.20, f=DISP, b=True, i=True)
    run(p, C.META["tagline"], DISP, 12.5, GOLD, b=True, i=True)
    tf = tbox(sl, M, sy2 + 1.18, 7.7, 0.22)
    p = para(tf, lead=LEAD_TIGHT, i=True)
    run(p, C.META["confidential"], SANS, 6.4, FAINT, i=True)
    return sl


def slide2(prs):
    sl = slide(prs, 2)
    y = header(sl, 1, "The proposition  ·  Who we are", "Every interaction is part of care.")
    footer(sl, 2)

    lw = 4.16
    gx = M + lw + 0.40
    gw = CW - lw - 0.40

    # narrative: both paragraphs in one frame so the renderer owns the flow
    nh = sum(h_of(t, lw, 7.7, lead=LEAD_BODY, pad=0) for t in C.PROPOSITION) + 0.16
    tf = tbox(sl, M, y, lw, nh)
    for t in C.PROPOSITION:
        p = para(tf, align=PP_ALIGN.JUSTIFY, after=5, lead=LEAD_BODY)
        run(p, t, SANS, 7.7, BODY)

    qy = y + nh + 0.12
    qh = h_of(C.PULLQUOTE, lw - 0.40, 9.6, f=DISP, i=True, lead=1.34) + 0.34
    rect(sl, M, qy, lw, qh, fill=CARD, line=LINE, lw=0.5, tag="quote")
    rect(sl, M, qy, 0.028, qh, fill=GOLD, tag="quoterule")
    tf = tbox(sl, M + 0.20, qy, lw - 0.36, qh, anchor=MSO_ANCHOR.MIDDLE)
    p = para(tf, lead=1.34, f=DISP, i=True)
    run(p, C.PULLQUOTE, DISP, 9.6, TAGCOL, i=True)

    # who we are, dark card, filled to the floor
    wy = qy + qh + 0.16
    wh = FLOOR - wy
    rect(sl, M, wy, lw, wh, fill=INK, tag="whocard")
    inner = lw - 0.44
    tf = tbox(sl, M + 0.22, wy + 0.19, inner, 0.18)
    p = para(tf, lead=LEAD_LABEL)
    run(p, "WHO WE ARE", SANS, 6.6, GOLD_L, b=True, track=2.2)
    hair(sl, M + 0.22, wy + 0.42, inner, color=GOLD)
    fh = h_of(C.WHO_WE_ARE, inner, 7.2, lead=LEAD_BODY)
    tf = tbox(sl, M + 0.22, wy + 0.56, inner, fh)
    p = para(tf, align=PP_ALIGN.JUSTIFY, lead=LEAD_BODY)
    run(p, C.WHO_WE_ARE, SANS, 7.2, ON_DARK)

    oy = wy + 0.56 + fh + 0.18
    hair(sl, M + 0.22, oy, inner, color=RULE_DK, h=0.006)
    tf = tbox(sl, M + 0.22, oy + 0.15, inner, 0.16)
    p = para(tf, lead=LEAD_LABEL)
    run(p, "WHAT IT SUPPORTS", SANS, 6.2, GOLD_L, b=True, track=1.9)
    oh = h_of(C.OUTCOME_STATEMENT, inner, 6.9, i=True, lead=LEAD_TIGHT)
    tf = tbox(sl, M + 0.22, oy + 0.38, inner, oh)
    p = para(tf, align=PP_ALIGN.JUSTIFY, lead=LEAD_TIGHT, i=True)
    run(p, C.OUTCOME_STATEMENT, SANS, 6.9, ON_DARK_MUTE, i=True)

    # capability grid, 3 x 3, heading + copy centred as one group
    block_label(sl, gx, y, gw, "What the programme builds")
    cols, rows_n, gap = 3, 3, 0.13
    cw = (gw - gap * (cols - 1)) / cols
    top = y + 0.32
    ch = (FLOOR - top - gap * (rows_n - 1)) / rows_n
    ti = cw - 0.32
    for i, (name, desc) in enumerate(C.PILLARS):
        cx = gx + (i % cols) * (cw + gap)
        cy = top + (i // cols) * (ch + gap)
        rect(sl, cx, cy, cw, ch, fill=CARD, line=LINE, lw=0.5, tag="pillar")
        rect(sl, cx, cy, cw, 0.026, fill=GOLD_L, tag="pillartop")
        nh2 = h_of(name, ti, 6.7, b=True, track=1.2, lead=LEAD_LABEL, pad=0)
        dh2 = h_of(desc, ti, 7.1, lead=LEAD_BODY, pad=0)
        gy = cy + (ch - (nh2 + 0.12 + dh2)) / 2
        tf = tbox(sl, cx + 0.16, gy, ti, nh2 + 0.02)
        p = para(tf, lead=LEAD_LABEL)
        run(p, name, SANS, 6.7, GOLD, b=True, track=1.2)
        tf = tbox(sl, cx + 0.16, gy + nh2 + 0.12, ti, dh2 + 0.02)
        p = para(tf, lead=LEAD_BODY)
        run(p, desc, SANS, 7.1, BODY)
    return sl


def slide3(prs):
    sl = slide(prs, 3)
    y = header(sl, 2, "Curriculum", "The twelve modules, and why in this order",
               dek=C.SEQUENCE_INTRO)
    footer(sl, 3)

    # built bottom-up: cohort strip, placements band, then the table takes the remainder
    coh_h = h_of(C.COHORT_NOTE, CW - 0.90, 6.6, lead=LEAD_TIGHT)
    coh_y = FLOOR - coh_h
    plc_body = max(h_of(t, 3.60, 7.0, lead=LEAD_TIGHT) for _, t in C.SEQUENCE_LOGIC)
    plc_h = 0.20 + plc_body
    plc_y = coh_y - 0.20 - plc_h
    cap_y = plc_y - 0.32

    widths = [3.02, 6.83, 1.02, 1.26]
    rows = [[m[0], m[1], m[2], m[3]] for m in C.MODULES]
    rows.append(["Total", "13 weeks, 3 sessions per week", "1–13", 39])
    head_h = 0.25
    avail = cap_y - 0.12 - y - head_h - 0.02
    row_h = min(0.276, (avail - 0.02) / 13.0)
    if row_h < 0.21:
        raise SystemExit(f"module table row height collapsed to {row_h:.3f}in")
    table(sl, M, y, widths, ["Module", "Focus", "Weeks", "Sessions"], rows,
          row_h=row_h, head_h=head_h, hsz=6.2, bsz=6.9, total=True,
          aligns=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.CENTER],
          col_sizes=[6.9, 6.7, 6.7, 6.9],
          col_colors=[INK, BODY, TAGCOL, INK], col_bold=[True, False, False, True])

    tf = tbox(sl, M, cap_y, CW, 0.20)
    p = para(tf, lead=LEAD_ROW)
    run(p, "ALLOCATION   ", SANS, 6.2, GOLD, b=True, track=1.6)
    run(p, C.ALLOCATION_NOTE, SANS, 6.8, BODY)

    gap = 0.26
    cw = (CW - gap * 2) / 3
    for i, (head, tail) in enumerate(C.SEQUENCE_LOGIC):
        cx = M + i * (cw + gap)
        if i:
            rect(sl, cx - gap / 2, plc_y + 0.02, 0.0075, plc_h - 0.04, fill=LINE, tag="vrule")
        tf = tbox(sl, cx, plc_y, cw, 0.18)
        p = para(tf, lead=LEAD_LABEL)
        run(p, BULLET + "  " + head.upper(), SANS, 6.4, GOLD, b=True, track=1.3)
        tf = tbox(sl, cx, plc_y + 0.21, cw, plc_body)
        p = para(tf, lead=LEAD_TIGHT)
        run(p, tail, SANS, 7.0, BODY)

    hair(sl, M, coh_y - 0.14, CW)
    tf = tbox(sl, M, coh_y, 0.85, 0.18)
    p = para(tf, lead=LEAD_ROW)
    run(p, "COHORTS", SANS, 6.2, GOLD, b=True, track=1.6)
    tf = tbox(sl, M + 0.90, coh_y, CW - 0.90, coh_h)
    p = para(tf, align=PP_ALIGN.JUSTIFY, lead=LEAD_TIGHT)
    run(p, C.COHORT_NOTE, SANS, 6.6, BODY)
    return sl


def slide4(prs):
    sl = slide(prs, 4)
    y = header(sl, 3, "Approach, delivery and measurement",
               "How it runs, and what the hospital receives")
    footer(sl, 4)

    # 1. how we teach, full width
    inner = CW - 0.40
    method = C.METHOD_NOTE + "  " + C.SESSION_ANATOMY
    mh = h_of(method, inner, 7.3, lead=LEAD_BODY) + 0.44
    rect(sl, M, y, CW, mh, fill=CARD, line=LINE, lw=0.5, tag="method")
    rect(sl, M, y, 0.028, mh, fill=GOLD, tag="methodrule")
    tf = tbox(sl, M + 0.20, y + 0.13, inner, 0.16)
    p = para(tf, lead=LEAD_LABEL)
    run(p, "HOW WE TEACH", SANS, 6.3, GOLD, b=True, track=1.8)
    tf = tbox(sl, M + 0.20, y + 0.33, inner, mh - 0.46)
    p = para(tf, align=PP_ALIGN.JUSTIFY, lead=LEAD_BODY)
    run(p, method, SANS, 7.3, BODY)

    # 2. role band at the foot
    rcols, rgap = 4, 0.24
    rw = (CW - 0.36 - rgap * (rcols - 1)) / rcols
    keep_h = max(h_of(k, rw, 6.2, lead=LEAD_TIGHT, pad=0) for _, _, k in C.ROLES)
    chip_h = 0.15 + 0.04 + keep_h
    band_h = 0.12 + 0.17 + 0.10 + chip_h * 2 + 0.14 + 0.13
    band_y = FLOOR - band_h
    rect(sl, M, band_y, CW, band_h, fill=HEAD, tag="roleband")
    tf = tbox(sl, M + 0.18, band_y + 0.12, 6.0, 0.17)
    p = para(tf, lead=LEAD_LABEL)
    run(p, "ALLOCATED BY ROLE  ·  WHAT EACH GROUP KEEPS", SANS, 6.4, GOLD, b=True, track=1.8)
    hair(sl, M + 0.18, band_y + 0.33, CW - 0.36, color=GOLD, h=0.008)
    ry0 = band_y + 0.44
    for i, (grp, _gain, keep) in enumerate(C.ROLES):
        rx = M + 0.18 + (i % rcols) * (rw + rgap)
        ry = ry0 + (i // rcols) * (chip_h + 0.14)
        tf = tbox(sl, rx, ry, rw, 0.15)
        p = para(tf, lead=LEAD_LABEL)
        run(p, grp, SANS, 6.6, INK, b=True)
        tf = tbox(sl, rx, ry + 0.19, rw, keep_h + 0.02)
        p = para(tf, lead=LEAD_TIGHT)
        run(p, keep, SANS, 6.2, BODY)

    # 3. five stages, horizontal
    sy = y + mh + 0.24
    sy = block_label(sl, M, sy, CW, "The five stages")
    scols, sgap = 5, 0.26
    scw = (CW - sgap * (scols - 1)) / scols
    sdh = max(h_of(d, scw, 6.6, lead=LEAD_TIGHT) for _, d in C.STAGES)
    for i, (name, desc) in enumerate(C.STAGES):
        cx = M + i * (scw + sgap)
        if i:
            rect(sl, cx - sgap / 2, sy, 0.0075, 0.20 + sdh, fill=LINE, tag="svrule")
        tf = tbox(sl, cx, sy - 0.03, 0.34, 0.24)
        p = para(tf, lead=LEAD_DISP, f=DISP, b=True)
        run(p, f"{i + 1:02d}", DISP, 11, GOLD_L, b=True)
        tf = tbox(sl, cx + 0.36, sy, scw - 0.36, 0.17)
        p = para(tf, lead=LEAD_LABEL)
        run(p, name, SANS, 6.7, INK, b=True, track=1.4)
        tf = tbox(sl, cx, sy + 0.22, scw, sdh)
        p = para(tf, lead=LEAD_TIGHT)
        run(p, desc, SANS, 6.6, BODY)

    # 4. two columns: deliverables | measurement
    top = sy + 0.22 + sdh + 0.26
    bottom = band_y - 0.22
    gap = 0.44
    cw = (CW - gap) / 2
    xs = [M, M + cw + gap]

    SIZES = [7.4, 7.2, 7.0, 6.8, 6.6, 6.4, 6.2]
    MG = 0.04

    x = xs[0]
    cy = block_label(sl, x, top, cw, "What the hospital receives")
    sz, ys, _ = fit_blocks(cy, bottom, lambda s: bullet_heights(C.DELIVERABLES, cw, s),
                           SIZES, min_gap=MG, tag="slide 4 deliverables")
    bullets_at(sl, x, ys, cw, C.DELIVERABLES, sz=sz, lead_bold=False)

    x = xs[1]
    cy = block_label(sl, x, top, cw, "How effectiveness is reported")
    dw = cw - 0.34

    def _rows(s):
        # the label run is bold and letter-spaced, so the line must be measured run-aware
        return [[(name + "   ", s - 0.4, SANS, True, False, 1.3),
                 (desc, s, SANS, False, False, 0.0)]
                for _, name, desc in C.MEASUREMENT_LEVELS]

    def _mh(s):
        hs = [max(0.15, ms.height_runs(r, dw, LEAD_TIGHT, pad=0.02)) for r in _rows(s)]
        hs.append(h_of(C.MEASUREMENT_CAVEAT, cw - 0.20, s - 0.2, i=True, lead=LEAD_TIGHT))
        return hs

    sz, ys, hs = fit_blocks(cy, bottom, _mh, SIZES, min_gap=MG, tag="slide 4 measurement")
    for (num, name, desc), ry, hh in zip(C.MEASUREMENT_LEVELS, ys, hs):
        tf = tbox(sl, x, ry - 0.02, 0.32, 0.22)
        p = para(tf, lead=LEAD_DISP, f=DISP, b=True)
        run(p, num, DISP, 10, GOLD_L, b=True)
        tf = tbox(sl, x + 0.34, ry, dw, hh)
        p = para(tf, lead=LEAD_TIGHT)
        run(p, name + "   ", SANS, sz - 0.4, INK, b=True, track=1.3)
        run(p, desc, SANS, sz, BODY)
    cav_h = hs[-1]
    cvy = ys[-1]
    rect(sl, x, cvy, 0.026, cav_h, fill=GOLD, tag="cavrule")
    tf = tbox(sl, x + 0.18, cvy, cw - 0.20, cav_h)
    p = para(tf, lead=LEAD_TIGHT, i=True)
    run(p, C.MEASUREMENT_CAVEAT, SANS, sz - 0.2, TAGCOL, i=True)
    return sl


def slide5(prs):
    sl = slide(prs, 5)
    y = header(sl, 4, "Commercial terms", "Rs 2,00,000 per month, for ninety days")
    footer(sl, 5)

    gap = 0.34
    cw = (CW - gap * 2) / 3
    xs = [M, M + cw + gap, M + 2 * (cw + gap)]

    # ── column 1: the fee, then the engagement at a glance
    x = xs[0]
    ci = cw - 0.40
    sub_h = h_of(C.FEE_HEADLINE_SUB, ci, 7.0, lead=LEAD_TIGHT)
    fee_h = 0.16 + 0.46 + sub_h + 0.07 + 0.20 + 0.16
    rect(sl, x, y, cw, fee_h, fill=INK, tag="feecard")
    tf = tbox(sl, x + 0.20, y + 0.15, ci, 0.46)
    p = para(tf, lead=LEAD_DISP, f=DISP, b=True)
    run(p, C.FEE_HEADLINE, DISP, 27, ON_DARK, b=True)
    tf = tbox(sl, x + 0.20, y + 0.63, ci, sub_h)
    p = para(tf, lead=LEAD_TIGHT)
    run(p, C.FEE_HEADLINE_SUB, SANS, 7.0, ON_DARK_MUTE)
    tf = tbox(sl, x + 0.20, y + 0.63 + sub_h + 0.07, ci, 0.20)
    p = para(tf, lead=LEAD_TIGHT)
    run(p, C.FEE_TOTAL_LINE, SANS, 7.0, GOLD_L, b=True)

    # the fee card above already carries fee, GST and engagement value, so the glance stays
    # operational rather than stating the money three times on one slide
    cy = block_label(sl, x, y + fee_h + 0.26, cw, "The engagement at a glance")
    glance = [r for r in C.GLANCE
              if r[0] not in ("PROFESSIONAL FEE", "GST", "TOTAL ENGAGEMENT VALUE")]
    c1_end = kv(sl, x, cy, cw, glance, label_w=1.52, lsz=5.9, vsz=7.1, min_h=0.32)

    # ── column 2: what the fee covers, then terms
    x = xs[1]
    cy = block_label(sl, x, y, cw, "What the fee covers")
    cy = accent_card(sl, x, cy, cw, C.FEE_SIMPLE, sz=7.6)
    cy = body(sl, x, cy + 0.18, cw, C.FEE_COVERS, sz=7.4)
    cy = block_label(sl, x, cy + 0.24, cw, "Terms")
    c2_end = bullets(sl, x, cy, cw, C.TERMS, sz=7.4, gap=0.13, lead_bold=False)

    # ── column 3: what we would need, then the mandatory note
    x = xs[2]
    cy = block_label(sl, x, y, cw, "What we would need from the hospital")
    cy = bullets(sl, x, cy, cw, C.NEEDS, sz=7.4, gap=0.14, lead_bold=False)
    c3_end = note_card(sl, x, cy + 0.26, cw, C.DISCLAIMER)

    # ── closing band, tall enough to close the slide off cleanly
    band_y = max(c1_end, c2_end, c3_end) + 0.30
    band_h = FLOOR + 0.02 - band_y
    if band_h < 0.90:
        raise SystemExit(f"slide 5 closing band collapsed to {band_h:.3f}in")
    rect(sl, M, band_y, CW, band_h, fill=INK, tag="closeband")
    tag_h = 0.30
    cl_h = h_of(C.CLOSING, 7.4, 7.1, i=True, lead=LEAD_TIGHT)
    grp = tag_h + 0.10 + cl_h
    gy = band_y + (band_h - grp) / 2
    tf = tbox(sl, M + 0.30, gy, 7.4, tag_h)
    p = para(tf, lead=1.20, f=DISP, b=True, i=True)
    run(p, C.META["tagline"], DISP, 13, GOLD_L, b=True, i=True)
    tf = tbox(sl, M + 0.30, gy + tag_h + 0.10, 7.4, cl_h)
    p = para(tf, lead=LEAD_TIGHT, i=True)
    run(p, C.CLOSING, SANS, 7.1, ON_DARK_MUTE, i=True)

    cxw = CW - 8.30
    tf = tbox(sl, M + 8.00, band_y, cxw, band_h, anchor=MSO_ANCHOR.MIDDLE)
    p = para(tf, align=PP_ALIGN.RIGHT, lead=1.10, f=DISP, b=True)
    run(p, C.META["contact_name"], DISP, 13, ON_DARK, b=True)
    p = para(tf, align=PP_ALIGN.RIGHT, lead=LEAD_LABEL)
    run(p, C.META["contact_role"], SANS, 6.6, GOLD_L, b=True, track=1.2)
    p = para(tf, align=PP_ALIGN.RIGHT, before=4, lead=LEAD_TIGHT)
    run(p, "+91 7411 099 183   ·   KAVITHA.PILLAI@LEVITATTE.COM", SANS, 6.2,
        ON_DARK_MUTE, track=0.4)
    p = para(tf, align=PP_ALIGN.RIGHT, lead=LEAD_TIGHT)
    run(p, "WWW.LEVITATTE.COM", SANS, 6.2, ON_DARK_MUTE, track=0.4)
    return sl


# ── geometry gate ──────────────────────────────────────────────────────────────
def guard():
    bad = []
    for page, tag, x, y, w, h in _SHAPES:
        if w <= 0 or h <= 0:
            bad.append(f"slide {page}: {tag} has non-positive size w={w:.3f} h={h:.3f}")
        if x < -0.001 or y < -0.001:
            bad.append(f"slide {page}: {tag} starts off-slide at x={x:.3f} y={y:.3f}")
        if x + w > W + 0.02 or y + h > H + 0.02:
            bad.append(f"slide {page}: {tag} runs off-slide to x2={x + w:.3f} y2={y + h:.3f}")
        if tag not in ("bg", "panel", "hair", "closeband") and y < FOOT_Y < y + h:
            bad.append(f"slide {page}: {tag} crosses the footer rule (y={y:.3f} h={h:.3f})")
    if bad:
        raise SystemExit("LAYOUT GUARD FAILED\n  " + "\n  ".join(bad))
    print(f"   guard: {len(_SHAPES)} shapes, all within bounds")


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "deck.pptx")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    for fn in (slide1, slide2, slide3, slide4, slide5):
        fn(prs)
    guard()
    prs.save(out)
    print(f"wrote {out}  ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
