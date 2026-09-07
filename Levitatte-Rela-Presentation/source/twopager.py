"""
Levitatte -> Rela: the two-page companion to the five-slide presentation.

Deliberately the less detailed edition. Same ivory / antique-gold system, same refined owl and
the same terms as the deck, condensed onto two sides of A4: a compact masthead, no cover, no
contents page, continuous flow, and the role outcomes reduced to one line each.

Typography is driven by BODY_PT / LINE so the whole document can be tightened in one place,
and build.sh asserts the result is exactly two pages.

    python3 twopager.py  Levitatte-Rela-2-Page-Proposal.docx
"""
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt

import content as C
import docx_brand as R

HERE = os.path.dirname(os.path.abspath(__file__))

# 7.9 / 1.18 is the largest setting that lands the document on exactly two sides of A4;
# build.sh re-checks the page count on every build rather than trusting this comment.
BODY_PT = float(os.environ.get("BODY_PT", "7.9"))
LINE = float(os.environ.get("LINE", "1.18"))
SMALL = BODY_PT - 0.5
CONTENT_W = 6.67          # A4 at 0.8in side margins


def setup(sec):
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.left_margin = sec.right_margin = Inches(0.80)
    sec.top_margin = Inches(0.62)
    sec.bottom_margin = Inches(0.58)
    sec.footer_distance = Inches(0.32)


def footer(sec, text):
    f = sec.footer
    p = f.paragraphs[0]
    p.text = ""
    R.spacing(p, before=0, after=0)
    R.para_border(p, "top", color=R.LINE_HEX, size=6, space=8)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(CONTENT_W), WD_TAB_ALIGNMENT.RIGHT)
    R.font(p.add_run(text), R.SANS, 6.6, R.MUTE, track=4)
    p.add_run("\t")
    R.add_page_field(p)


# ── masthead ───────────────────────────────────────────────────────────────────
def masthead(doc):
    if os.path.exists(R.OWL):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        R.spacing(p, before=0, after=1)
        p.add_run().add_picture(R.OWL, width=Inches(0.62))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    R.spacing(p, before=0, after=1)
    R.font(p.add_run("LEVITATTE"), R.DISP, 14.5, R.INK, bold=True, track=70)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    R.spacing(p, before=0, after=7)
    R.font(p.add_run("LEARNING  &  DEVELOPMENT"), R.SANS, 6.3, R.GOLD, bold=True, track=38)

    for line in C.META["title_lines"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        R.spacing(p, before=0, after=2, line=1.06)
        R.font(p.add_run(line), R.DISP, 14.5, R.INK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    R.spacing(p, before=3, after=2, line=1.16)
    R.font(p.add_run(C.META["subtitle"]), R.DISP, 9.5, R.TAGCOL, italic=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    R.spacing(p, before=0, after=6, line=1.24)
    R.font(p.add_run(C.META["descriptor"]), R.SANS, 6.3, R.GOLD, bold=True, track=26)

    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    R.table_hairlines(t, color=R.LINE_HEX, inner_h=False, inner_v=False, box=True)
    R.no_autofit(t, [Inches(5.05), Inches(1.62)])
    for i, c in enumerate(t.rows[0].cells):
        R.cell_margins(c, top=58, bottom=58, left=0 if i == 0 else 60, right=60)
    p = t.cell(0, 0).paragraphs[0]
    R.spacing(p, 0, 0, line=1.2)
    R.font(p.add_run(C.META["prepared_for"]), R.SANS, 8.0, R.INK)
    p = t.cell(0, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    R.spacing(p, 0, 0, line=1.2)
    R.font(p.add_run(C.META["date"]), R.SANS, 7.2, R.MUTE, bold=True, track=24)

    tail = doc.add_paragraph()
    R.spacing(tail, before=0, after=2)
    R.font(tail.add_run(""), R.SANS, 1, R.INK)


# ── section furniture ──────────────────────────────────────────────────────────
def head(doc, n, eyebrow, title):
    p = doc.add_paragraph()
    R.spacing(p, before=5, after=2)
    R.keep_with_next(p)
    R.font(p.add_run(f"{n:02d}   "), R.DISP, 8, R.GOLD_L, bold=True)
    R.font(p.add_run(eyebrow.upper()), R.SANS, 6.6, R.GOLD, bold=True, track=42)

    p = doc.add_paragraph()
    R.spacing(p, before=0, after=3.5, line=1.06)
    R.keep_with_next(p)
    R.font(p.add_run(title), R.DISP, 12.5, R.INK, bold=True)


def para(doc, text, size=BODY_PT, color=R.BODY, after=4.5, italic=False,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    R.spacing(p, before=0, after=after, line=LINE)
    p.alignment = align
    R.font(p.add_run(text), R.SANS, size, color, italic=italic)
    return p


def bullets(doc, items, size=BODY_PT, after=3.2, lead_bold=True):
    for it in items:
        p = doc.add_paragraph()
        R.spacing(p, before=0, after=after, line=LINE)
        pf = p.paragraph_format
        pf.left_indent = Inches(0.26)
        pf.first_line_indent = Inches(-0.26)
        pf.tab_stops.add_tab_stop(Inches(0.26), WD_TAB_ALIGNMENT.LEFT)
        R.font(p.add_run(R.BULLET), R.SANS, size, R.GOLD_L, bold=True)
        R.font(p.add_run("\t"), R.SANS, size, R.BODY)
        h, sep, tail = it.partition(": ")
        if sep and lead_bold and len(h) <= 62:
            R.font(p.add_run(h + ":"), R.SANS, size, R.INK, bold=True)
            R.font(p.add_run(" " + tail), R.SANS, size, R.BODY)
        else:
            R.font(p.add_run(it), R.SANS, size, R.BODY)


def callout(doc, text, size=BODY_PT, italic=True):
    t = doc.add_table(rows=1, cols=1)
    R.no_autofit(t, [Inches(CONTENT_W)])
    c = t.cell(0, 0)
    R.shade(c, R.CARD_HEX)
    R.cell_margins(c, top=90, bottom=90, left=130, right=130)
    R.table_hairlines(t, color=R.LINE_HEX, inner_h=False, inner_v=False, box=True)
    p = c.paragraphs[0]
    R.spacing(p, 0, 0, line=LINE)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    R.font(p.add_run(text), R.SANS, size, R.TAGCOL, italic=italic)
    sp = doc.add_paragraph()
    R.spacing(sp, before=0, after=2.5)
    R.font(sp.add_run(""), R.SANS, 1, R.INK)


def kv(doc, rows, label_w=1.72, lsz=6.4, vsz=None):
    vsz = vsz or BODY_PT
    t = doc.add_table(rows=len(rows), cols=2)
    R.table_hairlines(t, color=R.LINE_HEX, inner_h=True, inner_v=False, box=False)
    R.no_autofit(t, [Inches(label_w), Inches(CONTENT_W - label_w)])
    for i, (k, v) in enumerate(rows):
        for j, c in enumerate(t.rows[i].cells):
            R.cell_margins(c, top=36, bottom=36, left=0 if j == 0 else 40, right=40)
        p = t.cell(i, 0).paragraphs[0]
        R.spacing(p, 0, 0, line=1.14)
        R.font(p.add_run(k), R.SANS, lsz, R.GOLD, bold=True, track=22)
        p = t.cell(i, 1).paragraphs[0]
        R.spacing(p, 0, 0, line=1.14)
        R.font(p.add_run(v), R.SANS, vsz, R.INK)
    sp = doc.add_paragraph()
    R.spacing(sp, before=0, after=2.5)
    R.font(sp.add_run(""), R.SANS, 1, R.INK)


def module_table(doc):
    widths = [2.30, 3.22, 0.50, 0.65]
    cols = ["Module", "Focus", "Weeks", "Sessions"]
    rows = [[m[0], m[1], m[2], str(m[3])] for m in C.MODULES_SHORT]
    rows.append(["Total", "13 weeks, 3 sessions per week", "1–13", "39"])

    t = doc.add_table(rows=len(rows) + 1, cols=4)
    R.table_hairlines(t, color=R.LINE_HEX, inner_h=True, inner_v=False, box=False)
    R.no_autofit(t, [Inches(w) for w in widths])

    for j, name in enumerate(cols):
        c = t.cell(0, j)
        R.shade(c, R.HEAD_HEX)
        R.cell_margins(c, top=42, bottom=42, left=52 if j else 0, right=52)
        p = c.paragraphs[0]
        R.spacing(p, 0, 0, line=1.1)
        if j >= 2:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        R.font(p.add_run(name.upper()), R.SANS, 6.2, R.GOLD, bold=True, track=24)

    for i, row in enumerate(rows, start=1):
        is_total = i == len(rows)
        for j, val in enumerate(row):
            c = t.cell(i, j)
            if is_total:
                R.shade(c, R.HEAD_HEX)
            elif i % 2 == 0:
                R.shade(c, R.ZEBRA_HEX)
            R.cell_margins(c, top=34, bottom=34, left=52 if j else 0, right=52)
            p = c.paragraphs[0]
            R.spacing(p, 0, 0, line=1.1)
            if j >= 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            bold = is_total or j == 0
            color = R.INK if (is_total or j == 0) else (R.TAGCOL if j == 2 else R.BODY)
            R.font(p.add_run(val), R.SANS, SMALL if j == 1 else BODY_PT - 0.3,
                   color, bold=bold)
    sp = doc.add_paragraph()
    R.spacing(sp, before=0, after=2.5)
    R.font(sp.add_run(""), R.SANS, 1, R.INK)


def note(doc, text):
    p = doc.add_paragraph()
    R.spacing(p, before=5, after=0, line=LINE)
    R.para_border(p, "left", color="937425", size=14, space=8)
    p.paragraph_format.left_indent = Inches(0.11)
    R.font(p.add_run("NOTE   "), R.SANS, 6.3, R.GOLD, bold=True, track=30)
    R.font(p.add_run(text), R.SANS, SMALL, R.BODY, italic=True)


# ═══════════════════════════════════════════════════════════════════════════════
def build(out):
    doc = Document()
    R.set_page_background(doc)
    nm = doc.styles["Normal"]
    nm.font.name = R.SANS
    nm.font.size = Pt(BODY_PT)
    nm.font.color.rgb = R.BODY
    nm.paragraph_format.space_after = Pt(0)
    nm.paragraph_format.widow_control = True

    sec = doc.sections[0]
    setup(sec)
    footer(sec, C.META["footer"])
    masthead(doc)

    # 01 the proposition
    head(doc, 1, "The proposition", "Why this matters")
    for t in C.PROPOSITION:
        para(doc, t)
    para(doc, C.CAPABILITIES_LINE, size=SMALL, color=R.TAGCOL, after=4)

    # 02 commercial terms
    head(doc, 2, "Commercial terms", "The engagement at a glance")
    kv(doc, C.GLANCE)
    callout(doc, C.FEE_SIMPLE)
    para(doc, C.FEE_COVERS_BRIEF)

    # 03 curriculum
    head(doc, 3, "Curriculum", "The twelve modules, and why in this order")
    module_table(doc)
    bullets(doc, [f"{h}: {t}" for h, t in C.SEQUENCE_LOGIC_BRIEF], size=SMALL, after=2.2)

    # 04 outcomes by role
    head(doc, 4, "Outcomes by role", "What each group takes away")
    bullets(doc, [f"{g}: {gain}" for g, gain in C.ROLES_BRIEF], size=SMALL, after=2.2)

    # 05 approach and delivery
    head(doc, 5, "Approach and delivery", "How it runs, and what the hospital receives")
    bullets(doc, [f"{n}: {d}" for n, d in C.STAGES], size=SMALL, after=2.2)
    para(doc, C.METHOD_BRIEF, after=4)
    bullets(doc, C.DELIVERABLES, size=SMALL, after=2.2, lead_bold=False)
    para(doc, C.MEASUREMENT_BRIEF, size=SMALL, after=4)

    # 06 what we need, and in closing
    head(doc, 6, "What we need, and in closing", "Developing people who care for people")
    bullets(doc, C.NEEDS, size=SMALL, after=2.2, lead_bold=False)
    para(doc, " ".join(C.TERMS), size=SMALL, after=4)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    R.spacing(p, before=3, after=4, line=1.2)
    R.font(p.add_run(C.META["tagline"]), R.DISP, 11, R.GOLD, bold=True, italic=True)

    # run-in label rather than its own heading: the section label costs no extra line here
    p = doc.add_paragraph()
    R.spacing(p, before=0, after=4, line=LINE)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    R.font(p.add_run("Who we are.  "), R.SANS, BODY_PT, R.INK, bold=True)
    R.font(p.add_run(C.WHO_WE_ARE_BRIEF + " " + C.CLOSING), R.SANS, BODY_PT, R.BODY)

    p = doc.add_paragraph()
    R.spacing(p, before=2, after=0, line=1.2)
    R.font(p.add_run(C.META["contact_name"]), R.SANS, BODY_PT, R.INK, bold=True, track=10)
    R.font(p.add_run("   ·   " + C.META["contact_role"]), R.SANS, SMALL, R.BODY)
    p = doc.add_paragraph()
    R.spacing(p, before=1, after=0, line=1.2)
    R.font(p.add_run(C.META["contact_line"]), R.SANS, 6.6, R.MUTE, track=12)

    note(doc, C.DISCLAIMER)

    doc.save(out)
    print(f"wrote {out}  (body {BODY_PT}pt / line {LINE})")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "twopager.docx"))
