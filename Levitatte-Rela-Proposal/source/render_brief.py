"""
Concise (4-page) edition of the Levitatte -> Rela proposal.
Same ivory / antique-gold system and same refined owl as the full document, but with a compact
masthead instead of a cover, no contents page, continuous flow and tighter typography.
"""
import json
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt

import render as R

HERE = os.path.dirname(os.path.abspath(__file__))

TIGHT = os.environ.get("TIGHT") == "1"   # extra compression for the longer post-TDS edition
BODY_PT = 9.5
LINE = 1.25 if TIGHT else 1.32


def setup(sec):
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.left_margin = sec.right_margin = Inches(0.8)
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.68)
    sec.footer_distance = Inches(0.38)


def masthead(doc, meta):
    if os.path.exists(R.OWL):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        R.spacing(p, before=0, after=1)
        p.add_run().add_picture(R.OWL, width=Inches(0.84 if TIGHT else 0.98))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    R.spacing(p, before=0, after=1)
    R.font(p.add_run("LEVITATTE"), R.DISP, 16.5, R.INK, bold=True, track=76)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    R.spacing(p, before=0, after=11)
    R.font(p.add_run("LEARNING  &  DEVELOPMENT"), R.SANS, 6.5, R.GOLD, bold=True, track=40)

    for line in meta["title_lines"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        R.spacing(p, before=0, after=2, line=1.1)
        R.font(p.add_run(line), R.DISP, 16 if TIGHT else 17, R.INK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    R.spacing(p, before=5, after=3, line=1.25)
    R.font(p.add_run(meta["subtitle"]), R.DISP, 10.5, R.TAGCOL, italic=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    R.spacing(p, before=0, after=7, line=1.35)
    R.font(p.add_run(meta["descriptor"]), R.SANS, 6.5, R.GOLD, bold=True, track=28)

    # prepared-for / date strip on a single hairlined row
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    R.table_hairlines(t, color=R.LINE_HEX, inner_h=False, inner_v=False, box=True)
    R.no_autofit(t, [Inches(5.05), Inches(1.62)])
    for i, c in enumerate(t.rows[0].cells):
        R.cell_margins(c, top=80, bottom=80, left=0 if i == 0 else 60, right=60)
    p = t.cell(0, 0).paragraphs[0]
    R.spacing(p, 0, 0, line=1.25)
    R.font(p.add_run(meta["prepared_for"]), R.SANS, 8.5, R.INK)
    p = t.cell(0, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    R.spacing(p, 0, 0, line=1.25)
    R.font(p.add_run(meta["date"]), R.SANS, 7.5, R.MUTE, bold=True, track=26)

    tail = doc.add_paragraph()
    R.spacing(tail, before=0, after=10)
    R.font(tail.add_run(""), R.SANS, 1, R.INK)


def head(doc, eyebrow, title, n):
    p = doc.add_paragraph()
    R.spacing(p, before=8 if TIGHT else 10, after=3.5)
    R.keep_with_next(p)
    R.font(p.add_run(f"{n:02d}   "), R.DISP, 8.5, R.GOLD_L, bold=True)
    R.font(p.add_run(eyebrow.upper()), R.SANS, 7, R.GOLD, bold=True, track=44)

    p = doc.add_paragraph()
    R.spacing(p, before=0, after=6, line=1.12)
    R.keep_with_next(p)
    R.font(p.add_run(title), R.DISP, 14, R.INK, bold=True)


def para(doc, text, size=BODY_PT, color=R.BODY, after=5.5, italic=False):
    p = doc.add_paragraph()
    R.spacing(p, before=0, after=after, line=LINE)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    R.font(p.add_run(text), R.SANS, size, color, italic=italic)


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph()
        R.spacing(p, before=0, after=4, line=1.31)
        pf = p.paragraph_format
        pf.left_indent = Inches(0.28)
        pf.first_line_indent = Inches(-0.28)
        pf.tab_stops.add_tab_stop(Inches(0.28), WD_TAB_ALIGNMENT.LEFT)
        R.font(p.add_run(R.BULLET), R.SANS, BODY_PT, R.GOLD_L, bold=True)
        R.font(p.add_run("\t"), R.SANS, BODY_PT, R.BODY)
        h, sep, tail = it.partition(": ")
        if sep and len(h) <= 60:
            R.font(p.add_run(h + ":"), R.SANS, BODY_PT, R.INK, bold=True)
            R.font(p.add_run(" " + tail), R.SANS, BODY_PT, R.BODY)
        else:
            R.font(p.add_run(it), R.SANS, BODY_PT, R.BODY)


WIDTHS = {
    2: [Inches(1.85), Inches(4.82)],
    3: [Inches(1.42), Inches(2.72), Inches(2.53)],
    4: [Inches(2.02), Inches(3.25), Inches(0.66), Inches(0.74)],
}


def table(doc, columns, rows):
    n = len(columns)
    t = doc.add_table(rows=1, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    R.table_hairlines(t, inner_h=True, inner_v=False, box=False)
    for i, ctext in enumerate(columns):
        c = t.rows[0].cells[i]
        R.shade(c, R.HEAD_HEX)
        R.cell_margins(c, top=70, bottom=70, left=90, right=90)
        p = c.paragraphs[0]
        R.spacing(p, 0, 0)
        R.font(p.add_run(ctext.upper()), R.SANS, 6.5, R.GOLD, bold=True, track=26)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        total = str(row[0]).strip().lower() in ("total", "totals")
        for i, val in enumerate(row[:n]):
            c = cells[i]
            R.cell_margins(c, top=40 if TIGHT else 54, bottom=40 if TIGHT else 54, left=90, right=90)
            if total:
                R.shade(c, R.HEAD_HEX)
            elif ri % 2 == 1:
                R.shade(c, R.ZEBRA_HEX)
            p = c.paragraphs[0]
            R.spacing(p, 0, 0, line=1.2)
            R.font(p.add_run(str(val)), R.SANS, 8.5,
                   R.INK if (i == 0 or total) else R.BODY, bold=total or i == 0)
    # repeat the column header when a table continues onto the next page
    t.rows[0]._tr.get_or_add_trPr().append(R._el("w:tblHeader"))
    for row in t.rows:                      # never split a row mid-cell across a page break
        row._tr.get_or_add_trPr().append(R._el("w:cantSplit"))
    R.no_autofit(t, WIDTHS.get(n, [Inches(6.67 / n)] * n))
    tail = doc.add_paragraph()
    R.spacing(tail, before=0, after=6 if TIGHT else 8)
    R.font(tail.add_run(""), R.SANS, 1, R.INK)


def kv(doc, rows):
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    R.table_hairlines(t, inner_h=True, inner_v=False, box=False)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, c in enumerate(cells):
            R.cell_margins(c, top=62, bottom=62, left=0 if i == 0 else 90, right=90)
            if ri % 2 == 1:
                R.shade(c, R.ZEBRA_HEX)
        p = cells[0].paragraphs[0]
        R.spacing(p, 0, 0, line=1.2)
        R.font(p.add_run(str(row[0]).upper()), R.SANS, 6.5, R.MUTE, bold=True, track=24)
        p = cells[1].paragraphs[0]
        R.spacing(p, 0, 0, line=1.24)
        R.font(p.add_run(str(row[1])), R.SANS, 9, R.INK)
    R.no_autofit(t, WIDTHS[2])
    tail = doc.add_paragraph()
    R.spacing(tail, before=0, after=6 if TIGHT else 8)
    R.font(tail.add_run(""), R.SANS, 1, R.INK)


def pullquote(doc, text):
    p = doc.add_paragraph()
    R.spacing(p, before=5 if TIGHT else 7, after=6 if TIGHT else 8, line=1.3)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    R.font(p.add_run(text), R.DISP, 12, R.GOLD_L, bold=True, italic=True)


def callout(doc, text):
    p = doc.add_paragraph()
    R.spacing(p, before=6 if TIGHT else 8, after=0, line=1.28)
    p.paragraph_format.left_indent = Inches(0.16)
    R.para_border(p, "left", color="B08930", size=16, space=10)
    R.font(p.add_run("NOTE   "), R.SANS, 6.5, R.GOLD, bold=True, track=32)
    R.font(p.add_run(text), R.SANS, 8.5, R.BODY, italic=True)


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "brief_copy.json")
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, "proposal_brief.docx")
    data = json.load(open(src))
    meta, sections = data["meta"], data["sections"]

    doc = Document()
    R.set_page_background(doc)
    nm = doc.styles["Normal"]
    nm.font.name = R.SANS
    nm.font.size = Pt(BODY_PT)
    nm.font.color.rgb = R.BODY
    nm.paragraph_format.space_after = Pt(0)
    nm.paragraph_format.widow_control = True

    sec = doc.sections[0]
    setup(sec)
    R.build_footer(sec, meta["footer"])

    masthead(doc, meta)

    for i, s in enumerate(sections, 1):
        head(doc, s.get("eyebrow", ""), s["title"], i)
        if s.get("intro"):
            para(doc, s["intro"], size=10, color=R.TAGCOL, after=7)
        for b in s.get("blocks", []):
            k = b.get("type")
            if k == "para":
                para(doc, b["text"])
            elif k == "bullets":
                bullets(doc, b.get("items") or [])
            elif k == "table":
                table(doc, b["columns"], b["rows"])
            elif k == "kv":
                kv(doc, b["rows"])
            elif k == "pullquote":
                pullquote(doc, b["text"])
            elif k == "callout":
                callout(doc, b["text"])
            elif k == "signature":
                q = doc.add_paragraph()
                R.spacing(q, before=5 if TIGHT else 7, after=6 if TIGHT else 8)
                R.font(q.add_run(b["text"]), R.SANS, 8.5, R.INK, bold=True, track=12)

    if meta.get("contact_line"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        R.spacing(p, before=8, after=0, line=1.35)
        R.font(p.add_run(meta["contact_line"]), R.SANS, 6.5, R.MUTE, track=14)

    doc.save(out)
    print(f"wrote {out}  ({len(sections)} sections)")


if __name__ == "__main__":
    main()
